#!/usr/bin/env python3
"""Optional fixture regeneration. Requires python-pptx; never needed to build PPAM."""
from pathlib import Path
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

root = Path(__file__).resolve().parents[1]
deck = Presentation()
deck.slide_width = Inches(13.333333)
deck.slide_height = Inches(7.5)
for _ in range(4):
    deck.slides.add_slide(deck.slide_layouts[6])
data = CategoryChartData()
data.categories = ['A', 'B', 'C']
data.add_series('示例系列', (3, 5, 4))
shape = deck.slides[1].shapes.add_chart(
    XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(.42), Inches(4.375),
    Inches(6.25), Inches(2.7083), data)
shape.name = 'CHART'
deck.core_properties.author = 'Mac PowerPoint Batch Color Replacer contributors'
deck.core_properties.title = 'Synthetic acceptance fixture'
deck.save(root / 'templates/qa-fixture.pptx')
